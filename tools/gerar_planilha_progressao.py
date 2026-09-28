"""Planilha de progressao do Digilab (docs/TAREFA_PLANILHA_PROGRESSAO.md, Partes A e B).

    python tools/gerar_planilha_progressao.py              # escreve artifacts_progressao/ e docs/PROGRESSAO.md
    python tools/gerar_planilha_progressao.py --verificar  # gera duas vezes e confere: mesmas celulas, toda entrada com fonte

Le so o que ja existe: digilab/config.py, os JSONs de docs/, os artefatos (artifacts_treino/, artifacts_campanha/,
artifacts_combate_v21c/, experiments/configs/progressao/validacao_resultado.json) e, para registrar divergencias, as
constantes do jogo web (web/app/motor/constantes.ts). Nao roda luta nem vida: numero que nao existe em artefato fica
SEM FONTE, com a run que o produziria na aba `alertas`.

Toda celula numerica de entrada tem ao lado a fonte (arquivo + campo) e o status (CONFIG, ARTEFATO, REFERENCIA,
SEM FONTE, PROPOSTO). Valores derivados sao formulas do Excel. Fora de digilab/ (usa openpyxl e matplotlib).
"""
from __future__ import annotations

import argparse
import ast
import json
import re
import statistics
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from digilab.combate_v2.v21.catalogo import formas_da_especie, stats  # noqa: E402
from digilab.config import (  # noqa: E402
    CONFIG_BATERIA_V21, CONFIG_CAMPANHA, CONFIG_CATALOGO, CONFIG_COMBATE_V2C, CONFIG_COMBATE_V21C,
    CONFIG_ESCALA_LIGADA, CONFIG_EVOLUCAO_PESQUISA, CONFIG_SOMBRIO, CONFIG_TREINO,
)

SAIDA_DIR = RAIZ / "artifacts_progressao"
SAIDA_XLSX = SAIDA_DIR / "planilha_progressao.xlsx"
SAIDA_PNG = SAIDA_DIR / "curva_progressao.png"
SAIDA_MD = RAIZ / "docs" / "PROGRESSAO.md"

STATUS = ("CONFIG", "ARTEFATO", "REFERENCIA", "SEM FONTE", "PROPOSTO")
VERDE, AMARELO, VERMELHO = "C6EFCE", "FFEB9C", "FFC7CE"
CINZA, AZUL = "EDEDED", "DDEBF7"

# metas de taxa de vitoria (PROPOSTO, aba jornada/portoes): o momento em que o jogador PRECISA treinar
META_SEM_TREINO = 0.30          # pet do estagio sugerido, sem treino
META_COM_TREINO = 0.80          # com o treino tipico do estagio
FOLGA = 0.10                    # amarelo = ate 0,10 fora da meta

CFG = "digilab/config.py"
CRON = "docs/cronograma_interacoes.json"
PVE = "docs/pve_treino.json"
ARV = "docs/arvore_evolucao.json"
VAL = "experiments/configs/progressao/validacao_resultado.json"
RES_TREINO = "artifacts_treino/resumo.json"
CAMP_TREINO = "artifacts_campanha/treino.json"
GATES_V21C = "artifacts_combate_v21c/validacao/gates.json"
VIDAS = "artifacts_treino/vidas"
WEB = "web/app/motor/constantes.ts"
CALIB = "experiments/calibrar_progressao.py"


# ------------------------------------------------------------------ leitura das fontes

def _json(rel: str):
    return json.loads((RAIZ / rel).read_text(encoding="utf-8"))


def _ast_constante(rel: str, nome: str):
    """Valor literal de uma constante de modulo (sem importar o modulo, que rodaria o experimento)."""
    for no in ast.parse((RAIZ / rel).read_text(encoding="utf-8")).body:
        if isinstance(no, ast.Assign) and any(getattr(t, "id", None) == nome for t in no.targets):
            return ast.literal_eval(no.value)
    raise KeyError(f"{nome} nao encontrado em {rel}")


def _web_numero(padrao: str) -> float:
    m = re.search(padrao, (RAIZ / WEB).read_text(encoding="utf-8"))
    if not m:
        raise KeyError(f"padrao {padrao!r} nao encontrado em {WEB}")
    return float(m.group(1))


def _git(*args: str) -> str:
    return subprocess.run(["git", *args], cwd=RAIZ, capture_output=True, text=True, check=True).stdout.strip()


class Fontes:
    """Tudo o que a planilha le, carregado uma vez."""

    def __init__(self) -> None:
        self.cron = _json(CRON)
        self.pve = _json(PVE)
        self.arv = _json(ARV)
        self.val = _json(VAL)
        self.res_treino = _json(RES_TREINO)
        self.camp_treino = _json(CAMP_TREINO)
        self.gates = {g["gate"]: g for g in _json(GATES_V21C)["gates"]}
        self.alcance = _ast_constante(CALIB, "ALCANCE")
        self.niveis = _ast_constante(CALIB, "NIVEIS")
        self.vidas = self._ler_vidas()
        self.commit = _git("rev-parse", "--short", "HEAD")
        self.data_commit = _git("log", "-1", "--format=%cI")

    @staticmethod
    def _ler_vidas() -> dict:
        """Eventos do dono por vida (so contagem; nada e rejogado)."""
        out: dict[str, list[tuple[int, str]]] = {}
        for p in sorted((RAIZ / VIDAS).glob("*.jsonl")):
            evs = []
            for linha in p.read_text(encoding="utf-8").splitlines():
                e = json.loads(linha)
                if e["tipo"] == "acao_usuario":
                    evs.append((int(e["tick"]), str(e["valor"])))
            out[p.stem] = evs
        return out

    def interacoes_ate(self, braco: str, tick: int) -> float:
        n = [sum(1 for t, _ in evs if t < tick) for k, evs in self.vidas.items() if k.split("@")[0] == braco]
        return round(statistics.mean(n), 1)

    def sessoes_entre(self, braco: str, t0: int, t1: int) -> float:
        n = [sum(1 for t, v in evs if t0 <= t < t1 and v.startswith("treinar_")) for k, evs in self.vidas.items()
             if k.split("@")[0] == braco]
        return round(statistics.mean(n), 1)

    def taxa(self, npc: str, estagio: str, nivel: int) -> float:
        m = self.val["medidas"][npc][estagio]
        return round(statistics.mean(v[str(nivel)] for v in m.values()), 3)

    def nivel_tipico(self, lugar: int, estagio: str) -> int:
        teto = self.alcance[f"tier{lugar}"][estagio]
        return max(n for n in self.niveis[estagio] if n <= teto)


# ------------------------------------------------------------------ escrita com fonte e status

class Aba:
    """Uma aba com cabecalho; `entrada` escreve (valor, fonte, status) lado a lado e registra a celula."""

    def __init__(self, wb: Workbook, nome: str, cabecalho: list[str], larguras: dict[int, int] | None = None) -> None:
        self.ws = wb.create_sheet(nome)
        self.ws.append(cabecalho)
        for c in self.ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor=AZUL)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        self.ws.freeze_panes = "B2"
        for i in range(1, len(cabecalho) + 1):
            self.ws.column_dimensions[get_column_letter(i)].width = (larguras or {}).get(i, 16)
        self.linha = 1
        self.entradas: list[str] = []

    def nova_linha(self) -> int:
        self.linha += 1
        return self.linha

    def texto(self, col: int, valor) -> None:
        c = self.ws.cell(self.linha, col, valor)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    def entrada(self, col: int, valor, fonte: str, status: str) -> str:
        """valor numerico de entrada em `col`, fonte em col+1, status em col+2; devolve a referencia A1 do valor."""
        if status not in STATUS:
            raise ValueError(f"status desconhecido: {status}")
        if status != "SEM FONTE" and not fonte:
            raise ValueError("entrada sem fonte")
        c = self.ws.cell(self.linha, col, None if status == "SEM FONTE" else valor)
        if status == "SEM FONTE":
            c.value = "SEM FONTE"
            c.fill = PatternFill("solid", fgColor=CINZA)
        self.ws.cell(self.linha, col + 1, fonte).alignment = Alignment(wrap_text=True, vertical="top")
        s = self.ws.cell(self.linha, col + 2, status)
        if status == "PROPOSTO":
            s.font = Font(italic=True)
        ref = f"{get_column_letter(col)}{self.linha}"
        self.entradas.append(ref)
        return ref

    def formula(self, col: int, f: str, fmt: str | None = None) -> str:
        c = self.ws.cell(self.linha, col, f)
        if fmt:
            c.number_format = fmt
        return f"{get_column_letter(col)}{self.linha}"


def cores_status(ws, faixa: str, celula_topo: str) -> None:
    """Verde / amarelo / vermelho por texto da coluna de status (formula OK / PERTO / NAO FECHA)."""
    ws.conditional_formatting.add(faixa, FormulaRule(formula=[f'{celula_topo}="OK"'], fill=PatternFill("solid", fgColor=VERDE)))
    ws.conditional_formatting.add(faixa, FormulaRule(formula=[f'{celula_topo}="PERTO"'], fill=PatternFill("solid", fgColor=AMARELO)))
    ws.conditional_formatting.add(faixa, FormulaRule(formula=[f'{celula_topo}="NAO FECHA"'], fill=PatternFill("solid", fgColor=VERMELHO)))


def formula_status(sem: str, com: str, meta_sem: str, meta_com: str) -> str:
    """OK: sem treino <= meta e com treino >= meta; PERTO: ate FOLGA fora; senao NAO FECHA."""
    return (f'=IF(OR({sem}="SEM FONTE",{com}="SEM FONTE"),"SEM FONTE",IF(AND({sem}<={meta_sem},{com}>={meta_com}),"OK",'
            f'IF(AND({sem}<={meta_sem}+{FOLGA},{com}>={meta_com}-{FOLGA}),"PERTO","NAO FECHA")))')


# ------------------------------------------------------------------ abas

def aba_leia_me(wb: Workbook, f: Fontes) -> None:
    ws = wb.active
    ws.title = "leia_me"
    linhas = [
        ["Planilha de progressão do Digilab", ""],
        ["Tarefa", "docs/TAREFA_PLANILHA_PROGRESSAO.md (Partes A e B: mede e propõe; não muda regra)"],
        ["Commit das fontes", f.commit],
        ["Data do commit", f.data_commit],
        ["Como regenerar", "python tools/gerar_planilha_progressao.py   (--verificar gera duas vezes e compara)"],
        ["", ""],
        ["Abas", ""],
        ["referencias", "números dos jogos de referência (Tamagotchi, v-pets, Digimon World 1, Next Order), com a fonte"],
        ["jornada", "um marco por linha, do ovo ao endgame: tempo, interações, treino, poder esperado × exigido, taxas e metas"],
        ["estagios", "duração, intervalos das necessidades, janela, moeda, contadores, limiares, formas e fracasso"],
        ["portoes", "cada NPC da campanha e a arena: requisito, distância, taxa medida sem e com treino, meta"],
        ["economia_treino", "ganho, tetos, fadiga, forçar, sobrecarga, peso em combate; sessões e pontos por braço"],
        ["endgame", "o que existe depois do adulto: EXECUTADO / ESPECIFICADO / IDEIA"],
        ["principios", "os 7 princípios × Digilab: já tem / parcial / falta, e onde"],
        ["escala", "Parte B: escala atual × proposta × DW1, simulação nas metas, constantes absolutas do combate"],
        ["alertas", "onde a curva não fecha, divergências e células SEM FONTE (com a run que as produziria)"],
        ["", ""],
        ["Status das entradas", ""],
        ["CONFIG", "lido de digilab/config.py, de um JSON de docs/ ou de uma constante de experimento (fonte ao lado)"],
        ["ARTEFATO", "lido de um artefato em disco (fonte ao lado); contagens diretas de eventos das vidas contam aqui"],
        ["REFERENCIA", "número de um jogo de referência (guia de comunidade), só para desenho"],
        ["SEM FONTE", "o número não existe em nenhum artefato; a aba alertas diz qual run o produziria"],
        ["PROPOSTO", "meta ou modelo desta planilha (não é medida); precisa de aprovação"],
        ["", ""],
        ["Cores", ""],
        ["verde (OK)", f"dentro da meta: sem treino ≤ {META_SEM_TREINO} e com treino típico ≥ {META_COM_TREINO}"],
        ["amarelo (PERTO)", f"até {FOLGA} fora da meta"],
        ["vermelho (NAO FECHA)", "a curva não fecha: o conteúdo não exige o que deveria (ou exige demais)"],
        ["cinza", "SEM FONTE"],
        ["", ""],
        ["Modelo de poder (PROPOSTO)", "poder = total de stats × multiplicador de estágio. O total de stats do pet é o orçamento da "
         "espécie (60, soma zero) + peso_ponto_combate × pontos treinados; o do NPC é o total dele. O multiplicador de "
         "estágio é o que o combate aplica à vida e ao guts (ConfigCatalogo.mult_estagio). É um índice de ordem de "
         "grandeza: estilo e ε do NPC decidem mais que ele (ver alertas)."],
    ]
    for l in linhas:
        ws.append(l)
    ws["A1"].font = Font(bold=True, size=14)
    for r in (7, 18, 25):
        ws.cell(r, 1).font = Font(bold=True)
    for r, cor in ((26, VERDE), (27, AMARELO), (28, VERMELHO), (29, CINZA)):
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=cor)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


REFERENCIAS = [
    # jogo, item, valor, unidade, fonte
    ("Tamagotchi (1996)", "bebê → criança", 65, "min", "guias de comunidade (Thaao's Tamas, TamaVault, wiki do Tamagotchi)"),
    ("Tamagotchi (1996)", "adolescente aos", 3, "anos (1 ano = 1 dia real)", "idem"),
    ("Tamagotchi (1996)", "adulto aos", 6, "anos", "idem"),
    ("Tamagotchi (1996)", "erro de cuidado: chamado sem resposta em", 15, "min", "idem"),
    ("Tamagotchi (1996)", "criança 'boa' com no máximo", 1, "erro de cuidado", "idem"),
    ("Tamagotchi Uni", "criança e adolescente", 24, "h cada (aprox.)", "idem"),
    ("Tamagotchi Uni", "segunda chance: melhor adulto com 1–3 erros e pelo menos", 3, "passeios", "idem"),
    ("V-pets Digimon (Ver.20th/Color)", "estágio I", 10, "min", "guias do humulos"),
    ("V-pets Digimon (Ver.20th/Color)", "estágio II (6–12 h): mínimo", 6, "h", "idem"),
    ("V-pets Digimon (Ver.20th/Color)", "estágio III", 24, "h", "idem"),
    ("V-pets Digimon (Ver.20th/Color)", "estágio IV", 36, "h", "idem"),
    ("V-pets Digimon (Ver.20th/Color)", "estágio V", 48, "h", "idem"),
    ("Digimon World 1", "MetalGreymon: erros de cuidado até", 10, "erros", "guia de evolução do GameFAQs (SydMontague)"),
    ("Digimon World 1", "MetalGreymon: peso", 65, "g", "idem"),
    ("Digimon World 1", "MetalGreymon: bônus de batalhas", 30, "batalhas", "idem"),
    ("Digimon World 1", "vida: mínimo", 15, "dias de jogo", "idem"),
    ("Digimon World 1", "virar Ultimate soma", 96, "h de vida", "idem"),
    ("Digimon World 1", "In-Training: atributos até ~", 100, "pontos", "guia de evolução do GameFAQs; wiki Grindosaur"),
    ("Digimon World 1", "Rookie inicial: HP alvo", 1000, "HP", "idem"),
    ("Digimon World 1", "Rookie inicial: ataque alvo", 100, "pontos", "idem"),
    ("Digimon World 1", "primeiro Agumon selvagem: HP", 800, "HP", "idem"),
    ("Digimon World 1", "Champion 'sólido': HP", 1500, "HP", "idem"),
    ("Digimon World 1", "Champion: requisito de HP (mín.)", 1000, "HP", "idem"),
    ("Digimon World 1", "Champion: requisito de HP (máx.)", 2000, "HP", "idem"),
    ("Digimon World 1", "Ultimate: requisito de HP (mín.)", 3000, "HP", "idem"),
    ("Digimon World 1", "Ultimate: requisito de HP (máx.)", 7000, "HP", "idem"),
    ("Digimon World 1", "Ultimate: atributos (mín.)", 400, "pontos", "idem"),
    ("Digimon World 1", "Ultimate: atributos (máx.)", 600, "pontos", "idem"),
    ("Digimon World: Next Order", "bebê", 1, "h", "GameFAQs, PSNProfiles, Steam"),
    ("Digimon World: Next Order", "In-Training (mín.)", 3, "h", "idem"),
    ("Digimon World: Next Order", "Rookie (mín.)", 3, "dias", "idem"),
    ("Digimon World: Next Order", "Champion (mín.)", 5, "dias", "idem"),
    ("Digimon World: Next Order", "Ultimate (mín.)", 10, "dias", "idem"),
    ("Digimon World: Next Order", "Mega (mín.)", 14, "dias", "idem"),
    ("Digimon World: Next Order", "academia: ganho por hora de jogo (mín.)", 30, "pontos num stat", "idem"),
    ("Digimon World: Next Order", "batalhas: ganho por hora (mín.)", 100, "pontos em cada stat", "idem"),
    ("Digimon World: Next Order", "Ultimate no começo do jogo (mín.)", 1200, "por stat", "idem"),
    ("Digimon World: Next Order", "Ultimate no fim do jogo (mín.)", 7000, "por stat", "idem"),
    # jogos de fas recentes (referencia de mercado e de jornada; pedido do Matheus durante a Parte A)
    ("Digital Tamers: ReBorn (DragonRod)", "meta do ano: campeonatos (calendário como espinha)", 12, "campeonatos",
     "página do itch.io e do jogo; entrevista na Grindosaur"),
    ("Digital Tamers 2 (DragonRod, 2024)", "Jijimon chega à vila no nível", 3, "nível da vila", "idem"),
    ("Digital Tamers 2 (DragonRod, 2024)", "modo desafios abre com recrutados", 5, "Digimon recrutados (5+)", "idem"),
    ("Digital Tamers 2 (DragonRod, 2024)", "renascimento com herança: morrer no nível", 15, "nível (15+)", "idem"),
    ("Digital Tamers 2 (DragonRod, 2024)", "renascimento com herança: corações de amizade", 4, "corações (4+)", "idem"),
    ("Digital Tamers 2 (DragonRod, 2024)", "desafios", 120, "desafios (120+)", "idem"),
    ("Digital Tamers 2 (DragonRod, 2024)", "missões", 50, "missões (50+)", "idem"),
]

# o que os jogos de fas ensinam sem numero (aba referencias, parte de texto; princípios extras na aba principios)
REFERENCIAS_TEXTO = [
    ("Digital Tamers 2 / ReBorn", "luta 1x1 em plataforma: o jogador dá comandos, o Digimon luta; gratuito com doação; sem online "
     "(custo de servidor e trapaça); jogadores reclamam que o treino é complicado", "itch.io, página do jogo, Grindosaur, comentários"),
    ("Digital Monster Ver Web (Aderek, ago/2026)", "v-pet no navegador; o tempo continua com o navegador fechado; contas e save "
     "na nuvem; guia de campo das espécies encontradas; conjuntos de espécies por aparelho e da comunidade; app Android com "
     "widget e anúncios. Princípio declarado: a luta contribui para a progressão sem substituir o ciclo de cuidado",
     "página do jogo, comentários de jogadores"),
]


def aba_referencias(wb: Workbook) -> Aba:
    a = Aba(wb, "referencias", ["jogo", "item", "valor", "fonte", "status", "unidade"], {1: 30, 2: 52, 4: 60, 6: 26})
    for jogo, item, valor, un, fonte in REFERENCIAS:
        a.nova_linha()
        a.texto(1, jogo)
        a.texto(2, item)
        a.entrada(3, valor, f"{fonte} (via docs/TAREFA_PLANILHA_PROGRESSAO.md ou mensagem do Matheus)", "REFERENCIA")
        a.texto(6, un)
    a.nova_linha()
    for jogo, texto, fonte in REFERENCIAS_TEXTO:
        a.nova_linha()
        a.texto(1, jogo)
        a.texto(2, texto)
        a.texto(4, fonte)
        a.texto(5, "REFERENCIA (texto)")
    return a


def _npcs() -> list[tuple[str, dict]]:
    return list(CONFIG_CAMPANHA.npcs.items())


NOME_NPC = {"tier1_brutamontes": "Brutamontes (Gazimon)", "tier2_sentinela": "Sentinela (Tortomon)",
            "tier2_carcereiro": "Carcereiro", "tier2_porteiro": "Porteiro", "tier3_soberano": "Soberano"}
LUGAR = {1: "Campina", 2: "Torre de pedra", 3: "Cidadela"}


class Linha:
    """Cursor de colunas de uma linha: cada entrada ocupa (valor, fonte, status); texto e formula ocupam uma."""

    def __init__(self, a: Aba) -> None:
        self.a, self.col = a, 1

    def texto(self, v) -> None:
        self.a.texto(self.col, v)
        self.col += 1

    def entrada(self, v, fonte: str, status: str) -> str:
        ref = self.a.entrada(self.col, v, fonte, status)
        self.col += 3
        return ref

    def vazio3(self) -> None:
        self.a.texto(self.col, "—")
        self.col += 3

    def formula(self, f: str, fmt: str | None = None) -> str:
        ref = self.a.formula(self.col, f, fmt)
        self.col += 1
        return ref


def _cab(*itens) -> list[str]:
    """Cabecalho: nome simples = 1 coluna; ('nome', 3) = entrada (valor, fonte, status)."""
    out: list[str] = []
    for x in itens:
        out += [x[0], "fonte", "status"] if isinstance(x, tuple) else [x]
    return out


# B1: o que o jogador aprende em cada marco (no maximo duas mecanicas), como e ensinado e o sinal de que aprendeu.
# A voz do tutorial ate o adulto e o mentor original PROPOSTO (Bit, o zelador do laboratorio); do adulto em diante, o pet.
ENSINO = {
    "Ovo / eclosão": ("o pet é dele; o tempo do pet anda com o app aberto (fechar = pausa)",
                      "Bit, ao chocar: 'é seu; enquanto a janela estiver aberta, ele vive'. Uma frase por tela, some ao tocar",
                      "tocou no ovo e escolheu a velocidade do tempo"),
    "Fim do In-Training (bebê)": ("chamados e janela de 15 min (fome, higiene, carinho); atender rápido importa",
                                  "no primeiro chamado: o tempo cai para 1x, Bit aponta o balão e o botão certo; na revelação: "
                                  "'você atendeu X de Y chamados'",
                                  "atendeu 3 chamados dentro da janela"),
    "Rookie: Campina aberta": ("treino por botão (sem forçar); o primeiro rival pede treino",
                               "Bit leva à sala de treino no primeiro minuto do Rookie; no primeiro Desafiar perdido, "
                               "'faltou força: treine e volte'",
                               "treinou 3 sessões e venceu o Brutamontes"),
    "Fim do Rookie (criança)": ("ensinar: elogiar e repreender o que ele acabou de fazer; necessidade × vontade; a primeira recusa",
                                "na primeira ação espontânea: Bit mostra Elogiar/Repreender e o 'o que muda'; na primeira recusa, "
                                "o balão com o porquê",
                                "elogiou ou repreendeu 3 vezes na janela de 2 min"),
    "Champion: Torre de pedra": ("a linha que o caráter decidiu e a técnica; fadiga, 'Forçar' e o aviso de sobrecarga",
                                 "revelação do Champion mostra o eixo que decidiu; na fadiga 40 o pet pede pausa (Bit explica "
                                 "uma vez); no primeiro Forçar, a folha com o custo",
                                 "leu o porquê da linha; atendeu uma pausa (ou forçou sabendo o custo)"),
    "Fim do Champion (adolescente)": ("travar/vetar regras de combate; as 10 lutas de treino",
                                      "depois da primeira luta, o Resultado mostra a regra nova com Travar/Vetar; o contador "
                                      "'x de 10' aparece no topo",
                                      "travou ou vetou uma regra; completou as 10 lutas"),
    "Ultimate/Mega: Cidadela": ("forma final e passiva; a arena contra outros pets",
                                "o próprio pet explica a passiva na revelação (a partir do adulto ele fala)",
                                "desafiou o chefe"),
    "Fim da vida simulada (endgame)": ("o que fazer depois: arena, Melhor de 10, cenários; um novo ovo",
                                       "o pet sugere o que falta (formas vistas no Catálogo, chefe não vencido)",
                                       "voltou ao jogo depois de vencer o chefe"),
}


def aba_jornada(wb: Workbook, f: Fontes) -> Aba:
    """Um marco por linha. Entradas em trios (valor, fonte, status); derivados em formula."""
    def E(n: str) -> tuple:
        return (n, 3)
    cab = _cab("marco", "estágio", E("tempo desde o ovo (min)"), "tempo (h)",
               E("interações acumuladas — atencioso"), E("interações — aleatório"), E("interações — exigente"),
               E("pontos de treino típicos"), E("pontos medidos no fim da vida (atencioso)"), E("peso do ponto em combate"),
               E("total de stats do pet sem treino"), E("mult. de estágio do pet"),
               "poder esperado sem treino", "poder esperado com treino típico", "conteúdo liberado", "desafio",
               E("total de stats do desafio"), E("mult. de estágio do desafio"), "poder exigido", "distância (com treino − exigido)",
               E("taxa sem treino"), E("taxa com treino típico"), E("meta sem treino"), E("meta com treino"),
               "status da curva", "o que o jogador deve sentir",
               "o jogador aprende (B1)", "como é ensinado (tutorial contextual)", "sinal de que aprendeu")
    col_status = cab.index("status da curva") + 1
    col_sentir = cab.index("o que o jogador deve sentir") + 1
    a = Aba(wb, "jornada", cab, {1: 30, cab.index("conteúdo liberado") + 1: 30, cab.index("desafio") + 1: 30, col_sentir: 40,
                                 col_sentir + 1: 46, col_sentir + 2: 60, col_sentir + 3: 46})
    ev = CONFIG_EVOLUCAO_PESQUISA
    t_bebe = ev.dur_bebe_min
    t_cri = t_bebe + ev.dur_crianca_min
    t_ado = t_cri + ev.dur_adolescente_min
    fim_vida = f.res_treino["dias"] * 24 * 60
    npcs = dict(_npcs())
    torre = [k for k, v in npcs.items() if v["lugar"] == 2]
    tr = f.res_treino["por_braco"]["atencioso"]["treino_total"]["media"]
    marcos = [
        # marco, estagio, tick, fonte do tick, pontos tipicos, fonte, desafio, liberado, sentir
        ("Ovo / eclosão", "ovo", 0, f"{CRON} > _meta.eclosao (tick 0 = eclosão)", None, None, [],
         "o pet é dele; nome e espécie", "curiosidade: o que vai sair do ovo"),
        ("Fim do In-Training (bebê)", "bebe", t_bebe, f"{CFG} > ConfigEvolucaoPesquisa.dur_bebe_min", None, None, [],
         "ramo sereno/arisco (cuidado ≥ 80%)", "atender rápido mudou quem ele é"),
        ("Rookie: Campina aberta", "crianca", t_bebe, f"{CFG} > ConfigEvolucaoPesquisa.dur_bebe_min (a Campina abre com a criança)",
         f.alcance["tier1"]["crianca"], f"{CALIB} > ALCANCE['tier1']['crianca']", ["tier1_brutamontes"],
         "treino por botão; primeiro rival", "preciso treinar para vencer o primeiro rival"),
        ("Fim do Rookie (criança)", "crianca", t_cri, f"{CFG} > dur_bebe_min + dur_crianca_min", None, None, [],
         "linha pelo caráter + técnica", "o jeito que ensinei virou a linha dele"),
        ("Champion: Torre de pedra", "adolescente", t_cri, f"{CFG} > dur_bebe_min + dur_crianca_min",
         f.alcance["tier2"]["adolescente"], f"{CALIB} > ALCANCE['tier2']['adolescente']", torre,
         "3 rivais à escolha; forçar e sobrecarga", "um rival pede stats, outro pede estilo"),
        ("Fim do Champion (adolescente)", "adolescente", t_ado, f"{CFG} > soma das durações (+ 10 lutas de treino)",
         None, None, [], "forma adulta + passiva", "a forma final veio do caráter e do treino"),
        ("Ultimate/Mega: Cidadela", "adulto", t_ado, f"{CFG} > soma das durações",
         f.alcance["tier3"]["adulto"], f"{CALIB} > ALCANCE['tier3']['adulto']", ["tier3_soberano"],
         "chefe; arena", "o chefe exige tudo o que aprendi"),
        ("Fim da vida simulada (endgame)", "adulto", fim_vida, f"{RES_TREINO} > dias (14) × 1440", None, None, [],
         "arena, Melhor de 10, cenários", "meu pet é o meu, diferente dos outros"),
    ]
    for (nome, est, tick, fonte_t, pts, fonte_pts, desafio, liberado, sentir) in marcos:
        a.nova_linha()
        aprende, tutorial, sinal = ENSINO[nome]
        a.texto(col_sentir + 1, aprende)
        a.texto(col_sentir + 2, tutorial)
        a.texto(col_sentir + 3, sinal)
        L = Linha(a)
        L.texto(nome)
        L.texto(est)
        rt = L.entrada(tick, fonte_t, "ARTEFATO" if fonte_t.startswith("artifacts") else "CONFIG")
        L.formula(f"={rt}/60", "0.0")
        for braco in ("atencioso", "aleatorio", "exigente"):
            L.entrada(f.interacoes_ate(braco, tick),
                      f"{VIDAS}/{braco}@*.jsonl: eventos acao_usuario com tick < {tick} (média de 10 ovos)", "ARTEFATO")
        rp = None
        if pts is not None:
            rp = L.entrada(pts, fonte_pts, "CONFIG")
        else:
            L.vazio3()
        if nome.startswith("Fim da vida"):
            L.entrada(round(tr, 2), f"{RES_TREINO} > por_braco.atencioso.treino_total.media", "ARTEFATO")
        else:
            L.vazio3()
        rpeso = L.entrada(CONFIG_TREINO.peso_ponto_combate, f"{CFG} > ConfigTreino.peso_ponto_combate", "CONFIG")
        rtot = L.entrada(CONFIG_CATALOGO.orcamento_especie, f"{CFG} > ConfigCatalogo.orcamento_especie (soma zero dos modificadores)", "CONFIG")
        if est == "ovo":
            rmult = L.entrada(0, f"{ARV} > capacidade_por_estagio (o ovo não luta)", "CONFIG")
        else:
            rmult = L.entrada(CONFIG_CATALOGO.mult_estagio[est], f"{CFG} > ConfigCatalogo.mult_estagio['{est}']", "CONFIG")
        L.formula(f"={rtot}*{rmult}", "0.0")
        rcom = None
        if rp:
            rcom = L.formula(f"=({rtot}+{rpeso}*{rp})*{rmult}", "0.0")
        else:
            L.texto("—")
        L.texto(liberado)
        if not desafio:
            L.texto("—")
            a.texto(col_sentir, sentir)
            continue
        L.texto(" / ".join(NOME_NPC[k] for k in desafio) + (" (média dos rivais)" if len(desafio) > 1 else ""))
        tot = statistics.mean(npcs[k]["total_stats"] for k in desafio)
        rtn = L.entrada(tot, f"{CFG} > ConfigCampanha.npcs[{', '.join(desafio)}].total_stats" + (" (média)" if len(desafio) > 1 else ""), "CONFIG")
        est_npc = npcs[desafio[0]].get("estagio", "adulto")
        rmn = L.entrada(CONFIG_CATALOGO.mult_estagio[est_npc], f"{CFG} > mult_estagio['{est_npc}'] (estágio do NPC; sem o campo = adulto)", "CONFIG")
        rex = L.formula(f"={rtn}*{rmn}", "0.0")
        L.formula(f"={rcom}-{rex}", "0.0")
        lugar = npcs[desafio[0]]["lugar"]
        nv = f.nivel_tipico(lugar, est)
        sem = round(statistics.mean(f.taxa(k, est, 0) for k in desafio), 3)
        com = round(statistics.mean(f.taxa(k, est, nv) for k in desafio), 3)
        mais = " e dos rivais" if len(desafio) > 1 else ""
        rs = L.entrada(sem, f"{VAL} > medidas[{'/'.join(desafio)}][{est}][*]['0'] (média das formas{mais})", "ARTEFATO")
        rc = L.entrada(com, f"{VAL} > medidas[...][{est}][*]['{nv}'] (maior nível ≤ ALCANCE; média das formas{mais})", "ARTEFATO")
        ms = L.entrada(META_SEM_TREINO, "meta desta planilha (TAREFA_PLANILHA_PROGRESSAO B1: ~0,3 sem treino)", "PROPOSTO")
        mc = L.entrada(META_COM_TREINO, "meta desta planilha (B1: ~0,8 com treino típico)", "PROPOSTO")
        L.formula(formula_status(rs, rc, ms, mc))
        L.texto(sentir)
    c = get_column_letter(col_status)
    cores_status(a.ws, f"{c}2:{c}{a.linha}", f"{c}2")
    return a


def aba_estagios(wb: Workbook, f: Fontes) -> Aba:
    cab = ["estágio", "duração (min)", "fonte", "status",
           "intervalo fome (min)", "fonte", "status", "intervalo higiene (min)", "fonte", "status",
           "intervalo carinho (min)", "fonte", "status", "janela de chamado (min)", "fonte", "status",
           "sessões de treino/dia (alvo)", "fonte", "status", "teto de sessões/dia", "fonte", "status",
           "teto de pontos no estágio", "fonte", "status", "regras de combate", "fonte", "status",
           "mult. de estágio (vida e guts)", "fonte", "status",
           "moeda do estágio", "contadores zeram na evolução?", "limiares", "forma resultante", "forma de fracasso",
           "rota alternativa", "quem usa os intervalos"]
    a = Aba(wb, "estagios", cab, {1: 14, 32: 30, 34: 40, 35: 32, 36: 34, 37: 30, 38: 36})
    ev, ct, cv = CONFIG_EVOLUCAO_PESQUISA, CONFIG_TREINO, CONFIG_COMBATE_V21C
    durs = {"bebe": (ev.dur_bebe_min, "dur_bebe_min"), "crianca": (ev.dur_crianca_min, "dur_crianca_min"),
            "adolescente": (ev.dur_adolescente_min, "dur_adolescente_min"), "adulto": (None, "")}
    regras = {"bebe": 0, "crianca": cv.max_regras_crianca, "adolescente": cv.max_regras_adolescente, "adulto": cv.max_regras_adulto}
    textos = {
        "bebe": ("cuidado (fração de janelas atendidas)", "sim", f"cuidado ≥ {CONFIG_CATALOGO.limiar_cuidado_sereno} → sereno",
                 "Rookie sereno", "Rookie arisco (não é fim: é outra forma)", "falta (o cuidado do bebê decide uma vez)"),
        "crianca": ("caráter dominante", "sim", "eixo de caráter de maior valor", "Champion da linha (brigão, guardião, caçador, colosso)",
                    "sobrecarga ≥ 30: aviso (fala fria), sem mudar forma", "falta"),
        "adolescente": ("caráter complementar + 10 lutas de treino", "sim",
                        f"sinal do eixo complementar; {CONFIG_CATALOGO.sparring_minimo_adulto} lutas; sobrecarga ≥ {CONFIG_SOMBRIO.limiar_sombrio} → sombria",
                        "Ultimate/Mega (2 por linha)", "forma sombria da linha (forçar demais); sem 10 lutas, não evolui",
                        "falta (desempate por peso: cuidado_corporal.json não existe)"),
        "adulto": ("nenhuma (forma final); arena", "—", "—", "—", "—", "—"),
    }
    for est in ("bebe", "crianca", "adolescente", "adulto"):
        a.nova_linha()
        a.texto(1, est)
        d, campo = durs[est]
        if d is None:
            a.texto(2, "sem fim")
        else:
            a.entrada(2, d, f"{CFG} > ConfigEvolucaoPesquisa.{campo}", "CONFIG")
        iv = f.cron["estagios"][est]["intervalos_min"]
        for col, nec in ((5, "fome"), (8, "higiene"), (11, "carinho")):
            if est == "bebe":
                a.entrada(col, ev.chamados_bebe_min[nec], f"{CFG} > ConfigEvolucaoPesquisa.chamados_bebe_min['{nec}']", "CONFIG")
            else:
                a.entrada(col, iv[nec], f"{CRON} > estagios.{est}.intervalos_min.{nec} (nenhum código usa)", "CONFIG")
        a.entrada(14, ev.janela_chamado_min, f"{CFG} > ConfigEvolucaoPesquisa.janela_chamado_min", "CONFIG")
        if est == "bebe":
            a.texto(17, "não treina")
            a.texto(20, "—")
            a.texto(23, "—")
        else:
            a.entrada(17, ct.sessoes_alvo_dia[est], f"{CFG} > ConfigTreino.sessoes_alvo_dia['{est}']", "CONFIG")
            a.entrada(20, ct.teto_dia[est], f"{CFG} > ConfigTreino.teto_dia['{est}']", "CONFIG")
            a.entrada(23, ct.teto_por_estagio[est], f"{CFG} > ConfigTreino.teto_por_estagio['{est}']", "CONFIG")
        a.entrada(26, regras[est], f"{CFG} > ConfigCombateV21.max_regras_{est}" if est != "bebe" else f"{ARV} > capacidade_por_estagio.bebe (não luta)", "CONFIG")
        a.entrada(29, CONFIG_CATALOGO.mult_estagio[est], f"{CFG} > ConfigCatalogo.mult_estagio['{est}']", "CONFIG")
        moeda, zera, lim, forma, fracasso, rota = textos[est]
        a.texto(32, moeda)
        a.texto(33, zera)
        a.texto(34, lim)
        a.texto(35, forma)
        a.texto(36, fracasso)
        a.texto(37, rota)
        a.texto(38, "Python e web: só os chamados do bebê" if est == "bebe" else "ninguém (só no JSON)")
    return a


def aba_portoes(wb: Workbook, f: Fontes) -> Aba:
    cab = ["portão", "lugar", "estágio sugerido", "estilo",
           "total de stats", "fonte", "status", "ε (imprevisibilidade)", "fonte", "status",
           "taxa sem treino (média)", "fonte", "status", "taxa sem treino (melhor forma)", "fonte", "status",
           "taxa com treino típico (média)", "fonte", "status",
           "meta sem treino", "fonte", "status", "meta com treino", "fonte", "status",
           "status", "ganho do treino (com − sem)",
           "campanha: atencioso", "fonte", "status", "campanha: exigente", "fonte", "status",
           "exigente − atencioso (separa os donos?)", "nota"]
    a = Aba(wb, "portoes", cab, {1: 24, 4: 20, 37: 60})
    tv = f.val["todos_vencem"]
    linhas_status = []
    for k, v in _npcs():
        a.nova_linha()
        a.texto(1, NOME_NPC[k])
        a.texto(2, LUGAR[v["lugar"]])
        a.texto(3, v["estagio_minimo_sugerido"])
        a.texto(4, v["estilo"])
        a.entrada(5, v["total_stats"], f"{CFG} > ConfigCampanha.npcs['{k}'].total_stats", "CONFIG")
        a.entrada(8, v["epsilon"], f"{CFG} > ConfigCampanha.npcs['{k}'].epsilon", "CONFIG")
        rs = a.entrada(11, tv[k]["basico_media"], f"{VAL} > todos_vencem.{k}.basico_media", "ARTEFATO")
        a.entrada(14, tv[k]["basico_max"], f"{VAL} > todos_vencem.{k}.basico_max", "ARTEFATO")
        rc = a.entrada(17, tv[k]["treinado_media"], f"{VAL} > todos_vencem.{k}.treinado_media", "ARTEFATO")
        ms = a.entrada(20, META_SEM_TREINO, "meta desta planilha (B1)", "PROPOSTO")
        mc = a.entrada(23, META_COM_TREINO, "meta desta planilha (B1)", "PROPOSTO")
        linhas_status.append(a.formula(26, formula_status(rs, rc, ms, mc)))
        a.formula(27, f"={rc}-{rs}", "0.000")
        ra = a.entrada(28, round(f.camp_treino["por_braco"]["atencioso"][k]["media"], 3), f"{CAMP_TREINO} > por_braco.atencioso.{k}.media", "ARTEFATO")
        rx = a.entrada(31, round(f.camp_treino["por_braco"]["exigente"][k]["media"], 3), f"{CAMP_TREINO} > por_braco.exigente.{k}.media", "ARTEFATO")
        a.formula(34, f"={rx}-{ra}", "0.000")
        a.texto(35, "treino típico = ALCANCE do lugar (experiments/calibrar_progressao.py); NPCs imprevisíveis por decisão do "
                    "Matheus (2026-09-26): todo pet treinado vence, o pet sem treino vence com frequência")
    # arena
    a.nova_linha()
    a.texto(1, "Arena da pesquisa (V21C normalizada)")
    a.texto(2, "arena")
    a.texto(3, "adulto")
    a.texto(4, "pet × pet")
    ex = f.res_treino["exigente_x_atencioso"]
    a.texto(5, "—")
    a.texto(8, "—")
    a.entrada(11, round(ex["atencioso_contra_aleatorio"]["media"], 3), f"{RES_TREINO} > exigente_x_atencioso.atencioso_contra_aleatorio.media (não é 'sem treino': é atencioso × aleatório)", "ARTEFATO")
    a.texto(14, "—")
    a.entrada(17, round(ex["arena_exigente_contra_atencioso"]["media"], 3), f"{RES_TREINO} > exigente_x_atencioso.arena_exigente_contra_atencioso.media", "ARTEFATO")
    a.texto(26, "fora da meta (a arena mede o pet, não um portão)")
    a.texto(35, "a arena normaliza o total de stats ao orçamento: o treino muda o perfil, não o tamanho (pve_treino.json > normalizacao_pvp)")
    cores_status(a.ws, f"Z2:Z{a.linha}", "Z2")
    return a


def aba_economia(wb: Workbook, f: Fontes) -> Aba:
    cab = ["parâmetro", "valor", "fonte", "status", "nota"]
    a = Aba(wb, "economia_treino", cab, {1: 44, 3: 60, 5: 60})
    ct = CONFIG_TREINO
    params = [
        ("ganho por sessão (principal)", ct.ganho_principal, "ganho_principal", "× aptidão da espécie × (1 − pontos no stat/10) × (1 − fadiga/100)"),
        ("ganho por sessão (secundário)", ct.ganho_secundario, "ganho_secundario", ""),
        ("teto por stat", ct.teto_por_stat, "teto_por_stat", "retorno decrescente até aqui"),
        ("teto por estágio: Rookie", ct.teto_por_estagio["crianca"], "teto_por_estagio['crianca']", ""),
        ("teto por estágio: Champion", ct.teto_por_estagio["adolescente"], "teto_por_estagio['adolescente']", ""),
        ("teto por estágio: Ultimate/Mega", ct.teto_por_estagio["adulto"], "teto_por_estagio['adulto']", ""),
        ("teto na vida", ct.teto_vida, "teto_vida", ""),
        ("peso do ponto em combate", ct.peso_ponto_combate, "peso_ponto_combate", "1 → 3 (calibrar_progressao): cada ponto vale 3 de stat"),
        ("fadiga por sessão", ct.fadiga_por_sessao, "fadiga_por_sessao", ""),
        ("recupera fadiga por hora acordado", ct.recupera_fadiga_por_hora, "recupera_fadiga_por_hora", "zera ao dormir"),
        ("pede pausa com fadiga", ct.limiar_pausa, "limiar_pausa", ""),
        ("vira 'forçar' com fadiga", ct.limite_forcar, "limite_forcar", ""),
        ("intervalo mínimo entre sessões (min)", ct.intervalo_minimo_ticks, "intervalo_minimo_ticks", ""),
        ("forçar: multiplicador de ganho", ct.mult_forcar, "mult_forcar", ""),
        ("forçar: sobrecarga por sessão", ct.sobrecarga_forcar, "sobrecarga_forcar", "calibrado 15 → 25 (pve_treino.json diz 15)"),
        ("forçar: vínculo por sessão", ct.vinculo_forcar, "vinculo_forcar", ""),
        ("sobrecarga: aviso (fala fria)", CONFIG_SOMBRIO.limiar_aviso, "ConfigSombrio.limiar_aviso", ""),
        ("sobrecarga: forma sombria", CONFIG_SOMBRIO.limiar_sombrio, "ConfigSombrio.limiar_sombrio", "no momento da evolução para adulto"),
        ("sobrecarga: decai por hora bem cuidado", ct.decai_sobrecarga_hora_bem_cuidado, "decai_sobrecarga_hora_bem_cuidado", ""),
        ("sobrecarga: decai por sono completo", ct.decai_sobrecarga_sono_completo, "decai_sobrecarga_sono_completo", ""),
        ("cuidado por dia (orçamento)", ct.cuidado_por_dia, "cuidado_por_dia", "interações de cuidado + sessões-alvo do estágio"),
    ]
    for nome, v, campo, nota in params:
        a.nova_linha()
        a.texto(1, nome)
        a.entrada(2, v, f"{CFG} > {'ConfigTreino.' + campo if '.' not in campo else campo}", "CONFIG")
        a.texto(5, nota)
    a.nova_linha()
    a.texto(1, "orçamento de interações por vida (14 dias)")
    ro = a.entrada(2, f.res_treino["orcamento_por_vida"], f"{RES_TREINO} > orcamento_por_vida", "ARTEFATO")
    a.nova_linha()
    a.texto(1, "sessões de treino por vida (alvo)")
    rs = a.entrada(2, f.res_treino["slots_alvo"], f"{RES_TREINO} > slots_alvo", "ARTEFATO")
    a.nova_linha()
    a.texto(1, "custo do treino em interações (fração do orçamento)")
    a.formula(2, f"={rs}/{ro}", "0.00")
    a.texto(5, "fórmula: sessões-alvo / orçamento")
    # curva por braco e estagio
    ev = CONFIG_EVOLUCAO_PESQUISA
    t = [0, ev.dur_bebe_min, ev.dur_bebe_min + ev.dur_crianca_min, ev.dur_bebe_min + ev.dur_crianca_min + ev.dur_adolescente_min,
         f.res_treino["dias"] * 1440]
    a.nova_linha()
    a.nova_linha()
    for i, c in enumerate(["braço", "sessões no Rookie", "fonte", "status", "sessões no Champion", "fonte", "status",
                           "sessões no adulto", "fonte", "status", "pontos no fim da vida", "fonte", "status",
                           "forçadas por vida", "fonte", "status", "pontos acumulados por dia", "fonte", "status"], start=1):
        c_ = a.ws.cell(a.linha, i, c)
        c_.font = Font(bold=True)
        c_.fill = PatternFill("solid", fgColor=AZUL)
    for braco in ("atencioso", "exigente", "aleatorio", "relapso", "incoerente"):
        a.nova_linha()
        a.texto(1, braco)
        for col, (t0, t1, nome) in zip((2, 5, 8), ((t[1], t[2], "crianca"), (t[2], t[3], "adolescente"), (t[3], t[4], "adulto"))):
            a.entrada(col, f.sessoes_entre(braco, t0, t1), f"{VIDAS}/{braco}@*.jsonl: eventos treinar_* com {t0} ≤ tick < {t1} (média de 10)", "ARTEFATO")
        pb = f.res_treino["por_braco"][braco]
        a.entrada(11, round(pb["treino_total"]["media"], 2), f"{RES_TREINO} > por_braco.{braco}.treino_total.media", "ARTEFATO")
        a.entrada(14, pb["forcadas_media"], f"{RES_TREINO} > por_braco.{braco}.forcadas_media", "ARTEFATO")
        a.entrada(17, None, "", "SEM FONTE")
    return a


ENDGAME = [
    ("Arena contra bots (7 clássicos + 8 formas adultas)", "EXECUTADO", "web/app (aba Lutar · Arena)"),
    ("Melhor de 10", "EXECUTADO", "web/app/src/luta.ts > melhorDe10"),
    ("Campanha: Campina → Torre (3 rivais) → Cidadela", "EXECUTADO", "digilab/campanha.py; web/app (Lutar · Campanha)"),
    ("Formas sombrias (forçar o treino)", "EXECUTADO", "ConfigSombrio; combate_v2/v21; web"),
    ("Gambits híbridos (o pet aprende a lista; o dono trava/veta)", "EXECUTADO", "só na web (web/app/motor/gambits_pet.ts)"),
    ("Catálogo (galeria das 15 formas por brasão)", "EXECUTADO", "web/app (Diário · Catálogo)"),
    ("Cenários (pet de dono simulado vira rival fixo)", "EXECUTADO", "web/app (Diário · Cenários)"),
    ("Replay de luta no Diário", "EXECUTADO", "web/app (Resultado · Ver replay)"),
    ("Conversa por fala livre (LLM local)", "EXECUTADO", "digilab_fala_livre/, servidor/fala_livre.py"),
    ("Arena entre jogadores (pet × pet de outro dono)", "ESPECIFICADO", "pve_treino.json > normalizacao_pvp (sem rede)"),
    ("Modo produto: janela de 60 min, intervalos ×2, pausa (Cold Mode)", "ESPECIFICADO", "cronograma_interacoes.json > modo_produto (não implementado)"),
    ("Peso, superalimentação e marca 'robusto'", "ESPECIFICADO", "arvore_evolucao.json cita cuidado_corporal.json, que não existe"),
    ("Campeão secreto", "ESPECIFICADO", "docs/TAREFA_CAMPEAO_SECRETO.md (fora do git)"),
    ("Mapa de qualidade-diversidade (MAP-Elites) de pets", "ESPECIFICADO", "docs/TAREFA_QD_MAP_ELITES.md (fora do git)"),
    ("Segunda chance para quem errou cedo", "IDEIA", "princípio 4 (aba principios)"),
    ("Calendário de campeonatos (espinha do ano, como no Digital Tamers: ReBorn)", "IDEIA", "princípio E1"),
    ("Renascimento com herança (como no Digital Tamers 2)", "IDEIA", "princípio E2"),
    ("Dificuldade pelo ritmo de ganho (modo fácil)", "IDEIA", "princípio 6"),
]

PRINCIPIOS = [
    ("1. Uma moeda por estágio; contadores zeram na evolução", "já tem",
     "cuidado → ramo; caráter → linha; complementar + sparring → forma (arvore_evolucao.json > transicoes); "
     "contadores zeram (web/app/motor/motor.ts: sparring, chamados)"),
    ("2. Limiares grossos e legíveis (faixas)", "parcial",
     "limiares exatos (cuidado 0,8, eixo com sinal); a web mostra 'se evoluísse agora' com o porquê e as lutas x de 10"),
    ("3. Fracasso vira forma, não game over", "parcial",
     "arisco (pouco cuidado) e as formas sombrias (forçar); falta forma de fracasso para cuidado muito baixo no adulto"),
    ("4. Segunda chance para quem errou cedo", "falta", "o cuidado do bebê decide o ramo uma vez; nada recupera depois"),
    ("5. Indicador de proximidade da próxima forma", "parcial",
     "Evolução mostra a forma provável e o porquê; não há faixa Fix/H/M/L por requisito"),
    ("6. Dificuldade pelo ritmo de ganho, inimigos fixos", "falta",
     "NPCs fixos existem (ConfigCampanha); não há multiplicador de ganho por dificuldade (modo_produto só dobra intervalos)"),
    ("7. Treinar cedo rende mais", "falta",
     "ganho por sessão é o mesmo em todo estágio (ConfigTreino.ganho_principal); a proporção nem muda com a escala"),
]


# princípios extras, dos jogos de fas (Digital Tamers 2/ReBorn, Digital Monster Ver Web); pedido do Matheus
PRINCIPIOS_EXTRA = [
    ("E1. Calendário como espinha da jornada (ReBorn: 12 campeonatos no ano)", "falta",
     "a campanha é uma escada de 3 lugares, sem datas; o 'endgame' não tem meta no tempo"),
    ("E2. Renascimento com herança (DT2: morrer no nível 15+ com 4+ corações dá pontos para o próximo)", "falta",
     "um pet por vez; ovo novo apaga o anterior (web: Registro > Novo ovo); nada passa ao próximo pet"),
    ("E3. O tempo continua com o app fechado (Ver Web)", "falta",
     "web/app/src/tempo.ts: o tempo só anda com a página aberta (performance.now); fechar = pausa. Decisão de desenho a tomar"),
    ("E4. Guia de campo das espécies encontradas (Ver Web)", "já tem",
     "Diário · Catálogo: formas vistas × silhuetas, por brasão"),
    ("E5. A luta contribui sem substituir o ciclo de cuidado (Ver Web)", "parcial",
     "a luta ensina as regras e conta sparring; o cuidado decide o ramo e a obediência; mas a campanha hoje é vencida "
     "sem treino (alertas 1 e 2)"),
    ("E6. Treino simples de entender (reclamação em DT2)", "parcial",
     "um botão por exercício; fadiga, forçar, sobrecarga e aptidão somam 4 conceitos no mesmo painel"),
]


def aba_endgame(wb: Workbook) -> Aba:
    a = Aba(wb, "endgame", ["o que", "status", "onde"], {1: 60, 2: 16, 3: 70})
    for o, s, onde in ENDGAME:
        a.nova_linha()
        a.texto(1, o)
        a.texto(2, s)
        a.texto(3, onde)
        a.ws.cell(a.linha, 2).fill = PatternFill("solid", fgColor={"EXECUTADO": VERDE, "ESPECIFICADO": AMARELO, "IDEIA": CINZA}[s])
    return a


def aba_principios(wb: Workbook) -> Aba:
    a = Aba(wb, "principios", ["princípio", "Digilab", "onde"], {1: 52, 2: 14, 3: 100})
    for p, s, onde in PRINCIPIOS + PRINCIPIOS_EXTRA:
        a.nova_linha()
        a.texto(1, p)
        a.texto(2, s)
        a.texto(3, onde)
        a.ws.cell(a.linha, 2).fill = PatternFill("solid", fgColor={"já tem": VERDE, "parcial": AMARELO, "falta": VERMELHO}[s])
    return a


# ------------------------------------------------------------------ escala (Parte B)

# proposta de escala por estagio (PROPOSTO, B2): vida exibida alvo; atributos na mesma proporcao
ESCALA_PROPOSTA = dict(CONFIG_ESCALA_LIGADA.vida_por_estagio)
DW1 = {"bebe": (100, "In-Training: atributos até ~100 (HP de poucas centenas)"),
       "crianca": (1000, "Rookie inicial: HP alvo 1000"),
       "adolescente": (1500, "Champion 'sólido': HP ~1500 (requisitos 1000–2000)"),
       "adulto": (5000, "Ultimate: requisitos de HP 3000–7000 (meio da faixa)")}

# constantes absolutas do combate (Parte B, B2.3): o que nao escala junto com os stats
CONSTANTES_ABSOLUTAS = [
    ("dano do golpe rápido", "ConfigCombateV2.dano_golpe_rapido", CONFIG_COMBATE_V2C.dano_golpe_rapido, "dano", "vira relativo: × E_A / E_B (unidades de quem recebe)"),
    ("dano do disparo", "ConfigCombateV2.dano_disparo", CONFIG_COMBATE_V2C.dano_disparo, "dano", "idem"),
    ("dano do pesado", "ConfigCombateV2.dano_pesado", CONFIG_COMBATE_V2C.dano_pesado, "dano", "idem"),
    ("dano do parry (V2C)", "ConfigCombateV2C.dano_parry", CONFIG_COMBATE_V2C.dano_parry, "dano", "idem"),
    ("dano da investida", "ConfigCombateV21.dano_investida", CONFIG_COMBATE_V21C.dano_investida, "dano", "idem"),
    ("dano extra do duelista", "ConfigCombateV21.dano_extra_duelista", CONFIG_COMBATE_V21C.dano_extra_duelista, "dano", "idem"),
    ("vida base", "ConfigCombateV21C.hp_base", CONFIG_COMBATE_V21C.hp_base, "vida", "fica nas unidades próprias; a vida exibida é × E (não entra na luta)"),
    ("vida por vigor", "ConfigCombateV21C.hp_por_vigor", CONFIG_COMBATE_V21C.hp_por_vigor, "vida", "idem"),
    ("multiplicador de estágio (0,4 / 0,6 / 0,8 / 1,0) na vida e no guts", "ConfigCatalogo.mult_estagio", CONFIG_CATALOGO.mult_estagio["bebe"], "vida e guts", "sai da campanha: a razão E_A/E_B faz o papel dele na vida; no guts, decidir se fica (proposta: fica, é ritmo)"),
    ("regeneração de guts por turno (4 + 0,4 × recuperação)", "ConfigCombateV21.regen_base", CONFIG_COMBATE_V21C.regen_base, "guts", "fica absoluta (ritmo)"),
    ("custos de guts (golpe 10, pesado 15, disparo e esquiva 20)", "ConfigCombateV2.custo_*", CONFIG_COMBATE_V2C.custo_golpe_rapido, "guts", "guts não precisa escalar: é recurso de ritmo, não magnitude"),
    ("guts máximo, inicial, regen, carregar +25, defender +5", "ConfigCombateV21 / V2", CONFIG_COMBATE_V2C.guts_carregar, "guts", "idem (fica absoluto)"),
    ("limiares de guts nas condições (20, 40, 60)", "ConfigCombateV2.limiares_guts", 20, "guts", "idem"),
    ("multiplicadores de força/mira/guarda (0,7 + 0,03 × stat)", "ConfigCombateV21.mult_base / mult_por_ponto", CONFIG_COMBATE_V21C.mult_por_ponto, "stat", "dependem do stat de catálogo, que NÃO escala (o perfil fica)"),
    ("frações de vida nas condições (25/50/75%)", "ConfigCombateV2.fracoes_hp", 0.25, "fração", "já relativas"),
]


# Parte C: a escala aprovada esta em digilab/config.py > ConfigEscala (beta, vida por estagio, fatores calibrados)
BETA_TREINO = CONFIG_ESCALA_LIGADA.beta_treino
FATOR_NPC = dict(CONFIG_ESCALA_LIGADA.fator_npc)
ESCALA_VAL = "experiments/configs/escala/validacao.json"


def aba_escala(wb: Workbook, f: Fontes) -> Aba:
    a = Aba(wb, "escala", ["estágio", "vida exibida hoje (Brasa, sem treino)", "fonte", "status",
                           "stat de catálogo exibido hoje (vigor da Brasa)", "fonte", "status",
                           "vida proposta (escala E)", "fonte", "status", "DW1 (HP de referência)", "fonte", "status",
                           "crescimento hoje (× do bebê)", "crescimento proposto (× do bebê)", "crescimento DW1 (× do In-Training)",
                           "degrau hoje (× do estágio anterior)", "degrau proposto", "degrau DW1"], {1: 14})
    cv, cat = CONFIG_COMBATE_V21C, CONFIG_CATALOGO
    refs = {}
    ordem = ("bebe", "crianca", "adolescente", "adulto")
    for est in ordem:
        forma = next(x for x in formas_da_especie("brasa") if x.estagio == est)
        s_ = stats(forma)
        hp = (cv.hp_base + cv.hp_por_vigor * s_["vigor"]) * cat.mult_estagio[est]
        a.nova_linha()
        a.texto(1, est)
        rv = a.entrada(2, round(hp, 1), f"{CFG}: (hp_base + hp_por_vigor × vigor) × mult_estagio; forma {forma.rotulo} (catalogo.stats, V21C)", "CONFIG")
        rs = a.entrada(5, round(s_["vigor"], 2), f"digilab/combate_v2/v21/catalogo.py > stats({forma.rotulo})['vigor'] (pontos, soma 60)", "CONFIG")
        rp = a.entrada(8, ESCALA_PROPOSTA[est], f"{CFG} > ConfigEscala.vida_por_estagio['{est}'] (aprovada na Parte B)", "CONFIG")
        rd = a.entrada(11, DW1[est][0], f"DW1: {DW1[est][1]} (aba referencias)", "REFERENCIA")
        refs[est] = (rv, rs, rp, rd)
    for i, est in enumerate(ordem, start=2):
        rv, rs, rp, rd = refs[est]
        b = refs["bebe"]
        a.ws.cell(i, 14, f"={rv}/{b[0]}").number_format = "0.00"
        a.ws.cell(i, 15, f"={rp}/{b[2]}").number_format = "0.0"
        a.ws.cell(i, 16, f"={rd}/{b[3]}").number_format = "0.0"
        if i > 2:
            ant = refs[ordem[i - 3]]
            a.ws.cell(i, 17, f"={rv}/{ant[0]}").number_format = "0.00"
            a.ws.cell(i, 18, f"={rp}/{ant[2]}").number_format = "0.00"
            a.ws.cell(i, 19, f"={rd}/{ant[3]}").number_format = "0.00"
    esc = {est: refs[est][2] for est in ordem}

    def titulo(txt: str, colunas: list[str]) -> None:
        a.nova_linha()
        a.nova_linha()
        a.texto(1, txt)
        a.ws.cell(a.linha, 1).font = Font(bold=True, size=12)
        a.nova_linha()
        for k, c in enumerate(colunas, start=1):
            x = a.ws.cell(a.linha, k, c)
            x.font = Font(bold=True)
            x.fill = PatternFill("solid", fgColor=AZUL)
            x.alignment = Alignment(wrap_text=True, vertical="top")

    # --- como a escala entra no combate (B2.1 e B2.3)
    titulo("Desenho (PROPOSTO): perfil × escala", ["peça", "regra"])
    for peca, regra in [
        ("perfil", "os pontos do catálogo (espécie + temperamento + caráter + arquétipos, soma zero): o que a pesquisa compara; não muda"),
        ("escala E", "magnitude do lutador: E = E_estágio × (1 + β × pontos treinados). Vida exibida = vida de hoje sem o mult. de estágio × E"),
        ("combate", "cada lutador luta nas próprias unidades (vida e dano de hoje, sem o mult. de estágio); o dano que A causa em B "
                    "é multiplicado por E_A / E_B. Guts fica absoluto (é ritmo, não magnitude)"),
        ("invariância", "multiplicar E dos dois lados pelo mesmo fator dá E_A/E_B idêntico (mesmo float): a luta é exatamente a mesma. "
                        "Lutas entre pets do mesmo estágio e sem treino têm razão 1,0 exato: a bateria V21C não muda"),
        ("pesquisa", "a arena da pesquisa fixa E igual nos dois lados (adultos, treino normalizado ao orçamento): a escala não entra lá"),
        ("mult. de estágio de hoje", "sai do combate na campanha (a escala faz o papel dele); fica só na arena antiga e nas runs de pesquisa"),
    ]:
        a.nova_linha()
        a.texto(1, peca)
        a.texto(2, regra)

    # --- treino na escala (B2.5)
    titulo("Treino na escala (B2.5)", ["estágio", "E do estágio", "β por ponto", "fonte", "status", "pontos típicos no fim do estágio", "fonte", "status",
                                       "E com treino típico", "ganho em vida exibida por ponto", "ganho relativo por ponto"])
    beta_ref = None
    for est, pts, fonte in (("crianca", f.alcance["tier1"]["crianca"], "ALCANCE['tier1']['crianca']"),
                            ("adolescente", f.alcance["tier2"]["adolescente"], "ALCANCE['tier2']['adolescente']"),
                            ("adulto", f.alcance["tier3"]["adulto"], "ALCANCE['tier3']['adulto']")):
        a.nova_linha()
        a.texto(1, est)
        a.ws.cell(a.linha, 2, f"={esc[est]}")
        rb = a.entrada(3, BETA_TREINO, f"{CFG} > ConfigEscala.beta_treino (36 pontos → ×2,08)", "CONFIG")
        beta_ref = beta_ref or rb
        rpt = a.entrada(6, pts, f"{CALIB} > {fonte}", "CONFIG")
        a.formula(9, f"=B{a.linha}*(1+{rb}*{rpt})", "0")
        a.formula(10, f"=B{a.linha}*{rb}", "0.0")
        a.formula(11, f"={rb}/(1+{rb}*{rpt})", "0.0%")
    a.nova_linha()
    a.texto(1, "peso_ponto_combate")
    a.entrada(2, CONFIG_TREINO.peso_ponto_combate, f"{CFG} > ConfigTreino.peso_ponto_combate", "CONFIG")
    a.texto(5, "na nova base o ponto treinado vai para a escala (β) e para o perfil (proporções, como na arena normalizada); "
               "o ×3 absoluto sai da campanha (proposta: o perfil treinado é sempre normalizado ao orçamento)")

    # --- NPCs e simulação nas metas (B2.7, sem run)
    titulo("NPCs na escala e efeito nas metas (sem run)", [
        "portão", "E do pet sem treino", "E do pet com treino típico", "fator do NPC", "fonte", "status", "E do NPC",
        "razão sem treino (pet/NPC)", "razão com treino (pet/NPC)", "razão de hoje sem treino (poder)", "razão de hoje com treino (poder)",
        "taxa esperada sem treino", "fonte", "status", "taxa esperada com treino", "fonte", "status"])
    npcs = dict(_npcs())
    ct, cat = CONFIG_TREINO, CONFIG_CATALOGO
    razoes: list[tuple[str, float, float]] = []
    for k, v in npcs.items():
        lugar = v["lugar"]
        est = {1: "crianca", 2: "adolescente", 3: "adulto"}[lugar]
        est_npc = v.get("estagio", "adulto")
        pts = f.alcance[f"tier{lugar}"][est]
        a.nova_linha()
        a.texto(1, NOME_NPC[k])
        a.ws.cell(a.linha, 2, f"={esc[est]}")
        a.ws.cell(a.linha, 3, f"={esc[est]}*(1+{beta_ref}*{pts})").number_format = "0"
        rf = a.entrada(4, FATOR_NPC[k], f"{CFG} > ConfigEscala.fator_npc['{k}'] (calibrado: experiments/configs/escala/calibracao.json)", "CONFIG")
        a.formula(7, f"={esc[v['estagio_minimo_sugerido']]}*{rf}", "0")
        e_pet = ESCALA_PROPOSTA[est]
        e_npc = ESCALA_PROPOSTA[v["estagio_minimo_sugerido"]] * FATOR_NPC[k]
        razoes.append((NOME_NPC[k], e_pet / e_npc, e_pet * (1 + BETA_TREINO * pts) / e_npc))
        a.formula(8, f"=B{a.linha}/G{a.linha}", "0.00")
        a.formula(9, f"=C{a.linha}/G{a.linha}", "0.00")
        # conta explicita com os valores do config (orcamento, peso, mult. de estagio, total do NPC), para a origem ficar na celula
        me, mn, tot = cat.mult_estagio[est], cat.mult_estagio[est_npc], v["total_stats"]
        a.formula(10, f"={cat.orcamento_especie}*{me}/({tot}*{mn})", "0.00")
        a.formula(11, f"=({cat.orcamento_especie}+{ct.peso_ponto_combate}*{pts})*{me}/({tot}*{mn})", "0.00")
        val = json.loads((RAIZ / ESCALA_VAL).read_text(encoding="utf-8"))["resultado"][k] if (RAIZ / ESCALA_VAL).exists() else None
        if val:
            a.entrada(12, val["sem"], f"{ESCALA_VAL} > resultado.{k}.sem (sementes de validação)", "ARTEFATO")
            a.entrada(15, val["com"], f"{ESCALA_VAL} > resultado.{k}.com", "ARTEFATO")
        else:
            a.entrada(12, None, "", "SEM FONTE")
            a.entrada(15, None, "", "SEM FONTE")
    a.nova_linha()
    a.texto(1, "leitura")
    cruza = all(sem < 1 < com for _, sem, com in razoes)
    a.texto(2, ("com os fatores propostos a razão pet/NPC fica abaixo de 1,0 sem treino e acima com o treino típico em todos os portões ("
                if cruza else "atenção: com os fatores propostos a razão NÃO atravessa 1,0 em todos os portões (")
            + "; ".join(f"{n} {br(a_)} → {br(b_)}" for n, a_, b_ in razoes)
            + "). Os NPCs da Torre deixam de ter tamanho de adulto: passam à escala do Champion. As taxas vêm da validação da "
              "Parte C (experiments/calibrar_escala.py --validar); a Sentinela não fecha com nenhum fator (estilo tartaruga). As "
              "razões de hoje usam o modelo de poder da aba jornada (contas sobre os valores do config.py, visíveis na fórmula).")

    # --- E4 (B2.6)
    titulo("E4 redefinido (B2.6; definição final do Matheus depois da Parte C)", ["item", "antes da escala", "na escala"])
    for item, hoje, prop in [
        ("o que mede", "adulto × adolescente da mesma linha, mesma espécie e caráter, sem treino",
         "a evolução sempre importa: adulto SEM treino × adolescente sem treino e com o treino máximo do estágio (+22)"),
        ("meta", f"taxa do adulto em [{br(CONFIG_BATERIA_V21.taxa_min_e4)}, {br(CONFIG_BATERIA_V21.taxa_max_e4)}]",
         f"taxa do adulto ≥ {br(CONFIG_ESCALA_LIGADA.taxa_min_e4)} nos dois casos, sem teto (esmagar é de propósito na escala)"),
        ("por quê", "com escala de franquia o adulto sem treino esmaga o adolescente (E_adulto/E_Champion = 2,5)",
         "a Parte B propunha E4' (treino compensa: faixa [0,60, 0,85]) e E4'' (≥ 0,90); os dois reprovaram e o Matheus redefiniu"),
        ("resultado", ", ".join(f"{k} {br(v)}" for k, v in f.gates["E4_estagio_importa_sem_esmagar"]["valores"]["taxa_por_forma"].items()),
         e4_escala(f)),
    ]:
        a.nova_linha()
        a.texto(1, item)
        a.texto(2, hoje)
        a.texto(3, prop)

    # --- constantes absolutas (B2.3)
    titulo("Constantes absolutas do combate (B2.3)", ["constante", "campo", "valor", "fonte", "status", "tipo", "como tornar relativa"])
    for nome, campo, v, tipo, como in CONSTANTES_ABSOLUTAS:
        a.nova_linha()
        a.texto(1, nome)
        a.texto(2, campo)
        a.entrada(3, v, f"{CFG} > {campo}", "CONFIG")
        a.texto(6, tipo)
        a.texto(7, como)
    for col, w in ((1, 34), (2, 60), (3, 60), (7, 50)):
        a.ws.column_dimensions[get_column_letter(col)].width = w
    return a


def e4_escala(f: Fontes) -> str:
    """resultado do E4 na escala (artifacts_combate_v21c/validacao_escala/gates.json), se a run existir"""
    p = RAIZ / "artifacts_combate_v21c" / "validacao_escala" / "gates.json"
    if not p.exists():
        return "SEM FONTE (bateria V21C com a escala)"
    g = {x["gate"]: x for x in json.loads(p.read_text(encoding="utf-8"))["gates"]}.get("E4_escala_estagio_importa")
    return ("OK — " if g["ok"] else "REPROVA — ") + g["mensagem"] if g else "SEM FONTE (gate E4 na escala não encontrado)"


# ------------------------------------------------------------------ alertas e md

def alertas(f: Fontes) -> list[tuple[str, str, str, str, str]]:
    """(gravidade 1-5, titulo, numero, fonte, proposta). 1 = mais grave."""
    tv = f.val["todos_vencem"]
    sem = [tv[k]["basico_media"] for k, _ in _npcs()]
    e4 = f.gates["E4_estagio_importa_sem_esmagar"]["valores"]["taxa_por_forma"]
    difs = f.res_treino["exigente_x_atencioso"]["campanha_exigente_menos_atencioso"]
    web_car = _web_numero(r"escalaCarater:\s*([\d.]+)")
    brasa_bebe = stats(next(x for x in formas_da_especie("brasa") if x.estagio == "bebe"))["vigor"]
    brasa_ad = max(stats(x)["vigor"] for x in formas_da_especie("brasa") if x.estagio == "adulto")
    hp_b = (CONFIG_COMBATE_V21C.hp_base + CONFIG_COMBATE_V21C.hp_por_vigor * brasa_bebe) * CONFIG_CATALOGO.mult_estagio["bebe"]
    hp_a = (CONFIG_COMBATE_V21C.hp_base + CONFIG_COMBATE_V21C.hp_por_vigor * brasa_ad) * CONFIG_CATALOGO.mult_estagio["adulto"]
    return [
        ("1", "Campanha não exige treino: o pet sem treino do estágio sugerido vence os NPCs",
         f"{br(min(sem))}–{br(max(sem))} (média por NPC); com treino típico {br(min(tv[k]['treinado_media'] for k, _ in _npcs()))}–"
         f"{br(max(tv[k]['treinado_media'] for k, _ in _npcs()))}",
         f"{VAL} > todos_vencem.*.basico_media / treinado_media",
         "NPCs na escala proposta (aba escala) com distância alvo, e ε menor (0,2): o 'todo pet treinado vence' "
         "depende de ε 0,4–0,5, que também deixa o pet sem treino vencer. Recalibrar só depois da escala (Parte C)."),
        ("2", "A campanha não separa mais os donos",
         "exigente − atencioso: " + ", ".join(f"{k.split('_')[1]} {v['media']:+.3f}".replace(".", ",") for k, v in difs.items()),
         f"{RES_TREINO} > exigente_x_atencioso.campanha_exigente_menos_atencioso",
         "Consequência do alerta 1 (todo pet treinado vence 0,96–1,00). A pergunta 'forçar custa?' precisa de um portão "
         "que o pet sem treino perca; medir de novo depois da recalibração."),
        ("3", "E4: o adulto esmaga o adolescente da mesma linha",
         ", ".join(f"{k} {br(v)}" for k, v in sorted(e4.items(), key=lambda x: -x[1])) + f" (meta [{br(CONFIG_BATERIA_V21.taxa_min_e4)}, {br(CONFIG_BATERIA_V21.taxa_max_e4)}])",
         f"{GATES_V21C} > E4_estagio_importa_sem_esmagar.valores.taxa_por_forma",
         "Com uma escala de franquia a luta entre estágios fica desigual de propósito; redefinir o E4 (proposta na Parte B)."),
        ("4", "O pet não parece crescer: o número de catálogo quase não muda do bebê ao adulto",
         f"vigor da Brasa {brasa_bebe:.0f} no bebê → no máximo {brasa_ad:.0f} no adulto ({br(brasa_ad / brasa_bebe)}×; soma sempre 60); "
         f"a vida de combate vai de {hp_b:.0f} a {hp_a:.0f} ({br(hp_a / hp_b)}×)",
         "digilab/combate_v2/v21/catalogo.py > stats; ConfigCatalogo.mult_estagio; ConfigCombateV21C",
         "Separar perfil (pontos, soma zero) de escala (magnitude por estágio e treino); Parte B, aba escala."),
        ("5", "Durações dos estágios cortadas pela metade (em relação aos v-pets) limitam o treino por estágio",
         f"Rookie {CONFIG_EVOLUCAO_PESQUISA.dur_crianca_min / 60:.0f} h (v-pet 24 h), Champion {CONFIG_EVOLUCAO_PESQUISA.dur_adolescente_min / 60:.0f} h (v-pet 36 h); "
         f"sessões medidas no Rookie inteiro (atencioso): {br(f.sessoes_entre('atencioso', CONFIG_EVOLUCAO_PESQUISA.dur_bebe_min, CONFIG_EVOLUCAO_PESQUISA.dur_bebe_min + CONFIG_EVOLUCAO_PESQUISA.dur_crianca_min), 1)}, "
         f"com alvo de {CONFIG_TREINO.sessoes_alvo_dia['crianca']} por dia e teto de {CONFIG_TREINO.teto_por_estagio['crianca']:g} pontos no estágio",
         f"{CRON} > estagios.*.nota; {VIDAS} (contagem de treinar_*)",
         "Com o Rookie de 12 h (parte dele dormindo) cabem poucas sessões; os pontos no fim do Rookie são SEM FONTE (ver "
         "abaixo). Se ficarem longe do teto: o estágio cresce, o teto desce, ou o ganho por sessão sobe no Rookie (princípio 7)."),
        ("6", "Divergência: NPCs de docs/pve_treino.json × config.py",
         "JSON: tier 1 72 agressivo, tier 2 80 contragolpe, tier 3 90 adaptativo; config: Brutamontes 96, Torre 72/90/90, Soberano 110",
         f"{PVE} > npcs; {CFG} > ConfigCampanha.npcs",
         "Registrado em REGISTRO_DESVIOS (escada rebalanceada, 2026-09-25/26). Atualizar o JSON ou marcá-lo como histórico."),
        ("6", "Divergência: sobrecarga por sessão forçada",
         f"JSON 15 × config {CONFIG_TREINO.sobrecarga_forcar:g}", f"{PVE} > treino.forcar.custos.sobrecarga; {CFG} > ConfigTreino.sobrecarga_forcar",
         "Calibração registrada (15 → 25). Atualizar o JSON."),
        ("6", "Divergência: cuidado_corporal.json citado e ausente",
         "marca 'robusto' e desempate da forma adulta por peso", f"{ARV} > transicoes (marca_robusto, desempate)",
         "Especificar o cuidado corporal ou tirar a referência da árvore."),
        ("6", "Divergência: intervalos das necessidades da criança em diante e modo produto só no JSON",
         "criança 90/120/90, adolescente 150/120/120, adulto 180/120/180; modo produto janela 60, intervalos ×2",
         f"{CRON} > estagios.*.intervalos_min, modo_produto",
         "Nenhum código usa (Python nem web): só o bebê chama. Decidir se a criança chama ou tirar do JSON."),
        ("6", "Divergência: escala do caráter pelos traços na web",
         f"web {web_car:g} × Python escala_carater {CONFIG_EVOLUCAO_PESQUISA.escala_carater:g}",
         f"{WEB} > CAT.escalaCarater; {CFG} > ConfigEvolucaoPesquisa.escala_carater",
         "O 0,1 é do caráter PROVISÓRIO pelos traços (pets v1/v2, só na web); os pets v3 usam o desvio da política. Registrar ou alinhar."),
        ("6", "Divergência: o bebê luta na web (×0,4) e não no Python; a condição física vem da hora da luta na web",
         "—", "REGISTRO_DESVIOS (2026-09-24 e 2026-09-25)", "Já registrado; manter."),
        ("6", "Divergência: tempo da web",
         "velocidade padrão 10 min/s; com chamado aberto cai para 1x", "web/app/src/tempo.ts > VELOCIDADES, velocidade",
         "O Python simula 14 dias de minuto a minuto; a web deixa o jogador acelerar. A jornada em horas reais depende disso."),
        ("7", "Células SEM FONTE: pontos de treino acumulados por dia (aba economia_treino)",
         "5 células (uma por braço)", f"{VIDAS}/*.treino.json (só o total da vida)",
         "Run: rejogar as 50 vidas de artifacts_treino/vidas com MotorTreino gravando os pontos a cada dia "
         "(~1–2 min; leitura das vidas já gravadas, sem vida nova). Não rodada: a tarefa proíbe experimento para preencher."),
        ("7", "Célula SEM FONTE: taxa de vitória do pet típico contra os bots da arena do jogo",
         "a arena do jogo (bots, stats normalizados) não tem artefato por estágio", "—",
         "Run: pets de artifacts_treino por estágio × 15 bots da arena, 100 partidas (~5 min com 8 processos)."),
    ]


def aba_alertas(wb: Workbook, f: Fontes) -> Aba:
    a = Aba(wb, "alertas", ["gravidade (1 = mais grave)", "alerta", "número", "fonte", "proposta (só proposta)"], {1: 12, 2: 50, 3: 60, 4: 60, 5: 80})
    for g, t, n, fo, p in alertas(f):
        a.nova_linha()
        a.texto(1, int(g))
        a.texto(2, t)
        a.texto(3, n)
        a.texto(4, fo)
        a.texto(5, p)
        cor = VERMELHO if int(g) <= 3 else AMARELO if int(g) <= 5 else CINZA
        a.ws.cell(a.linha, 1).fill = PatternFill("solid", fgColor=cor)
    return a


def curva_png(f: Fontes, destino: Path) -> dict:
    """Poder esperado (sem e com treino tipico) x exigido nos tres lugares da campanha, e as taxas medidas."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    npcs = dict(_npcs())
    ct, cat = CONFIG_TREINO, CONFIG_CATALOGO
    pontos = []
    for lugar, est in ((1, "crianca"), (2, "adolescente"), (3, "adulto")):
        ks = [k for k, v in npcs.items() if v["lugar"] == lugar]
        nv = f.nivel_tipico(lugar, est)
        tot_npc = statistics.mean(npcs[k]["total_stats"] for k in ks)
        m_npc = cat.mult_estagio[npcs[ks[0]].get("estagio", "adulto")]
        pontos.append({
            "lugar": LUGAR[lugar], "estagio": est, "nivel": nv,
            "sem": cat.orcamento_especie * cat.mult_estagio[est],
            "com": (cat.orcamento_especie + ct.peso_ponto_combate * f.alcance[f"tier{lugar}"][est]) * cat.mult_estagio[est],
            "exigido": tot_npc * m_npc,
            "taxa_sem": statistics.mean(f.taxa(k, est, 0) for k in ks),
            "taxa_com": statistics.mean(f.taxa(k, est, nv) for k in ks),
        })
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(11, 4.2))
    x = range(len(pontos))
    rot = [f"{p['lugar']}\n({p['estagio']})" for p in pontos]
    ax1.plot(x, [p["exigido"] for p in pontos], "o-", color="#B0213F", label="poder exigido (NPC)")
    ax1.plot(x, [p["sem"] for p in pontos], "s--", color="#8A9AC0", label="pet sem treino")
    ax1.plot(x, [p["com"] for p in pontos], "s-", color="#2C73E0", label="pet com treino típico")
    ax1.set_xticks(list(x), rot)
    ax1.set_ylabel("total de stats × mult. de estágio (PROPOSTO)")
    ax1.set_title("Poder esperado × exigido")
    ax1.legend(fontsize=8)
    ax2.plot(x, [p["taxa_sem"] for p in pontos], "s--", color="#8A9AC0", label="sem treino (medido)")
    ax2.plot(x, [p["taxa_com"] for p in pontos], "s-", color="#2C73E0", label="com treino típico (medido)")
    ax2.axhline(META_SEM_TREINO, color="#8A9AC0", lw=.8, ls=":")
    ax2.axhline(META_COM_TREINO, color="#2C73E0", lw=.8, ls=":")
    ax2.set_xticks(list(x), rot)
    ax2.set_ylim(0, 1.05)
    ax2.set_ylabel("taxa de vitória contra o NPC")
    ax2.set_title("Taxa medida × meta (pontilhado, PROPOSTO)")
    ax2.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(destino, dpi=110, metadata={"Software": None, "Creation Time": None})
    plt.close(fig)
    return {"pontos": pontos}


def br(x: float, d: int = 2) -> str:
    return f"{x:.{d}f}".replace(".", ",")


def leitura(ps: list[dict]) -> str:
    """Frase do resumo tirada dos numeros (nao escrita a mao)."""
    abaixo = [p["lugar"] for p in ps if p["com"] < p["exigido"]]
    acima = [p["lugar"] for p in ps if p["com"] >= p["exigido"]]
    sem = [p["taxa_sem"] for p in ps]
    partes = []
    if abaixo:
        partes.append(f"pelo modelo de poder o pet com treino típico fica abaixo do NPC em {', '.join(abaixo)} e mesmo assim vence "
                      + ", ".join(br(p['taxa_com']) for p in ps if p['lugar'] in abaixo))
    if acima:
        partes.append(f"fica acima em {', '.join(acima)}")
    return ("; ".join(partes) + f". O pet sem treino fica abaixo do NPC nos três lugares e vence {br(min(sem))}–{br(max(sem))}: "
            "o estilo e a imprevisibilidade (ε 0,4–0,5) dos NPCs decidem mais que o tamanho. O momento em que o jogador "
            "**precisa** treinar não existe hoje.")


def escrever_md(f: Fontes, curva: dict) -> None:
    al = alertas(f)
    graves = [x for x in al if int(x[0]) <= 5][:5]
    ps = curva["pontos"]
    linhas = [
        "# Progressão do Digilab (resumo)",
        "",
        f"Gerado por `tools/gerar_planilha_progressao.py` (commit das fontes `{f.commit}`). A planilha completa, com a fonte de "
        "cada número, é `artifacts_progressao/planilha_progressao.xlsx`. Metas marcadas PROPOSTO dependem de aprovação.",
        "",
        "## A curva",
        "",
        "![poder esperado × exigido](../artifacts_progressao/curva_progressao.png)",
        "",
        "| lugar | estágio | poder do pet sem / com treino | poder exigido | taxa sem treino | taxa com treino típico | meta (PROPOSTO) |",
        "|---|---|---|---|---|---|---|",
    ]
    for p in ps:
        linhas.append(f"| {p['lugar']} | {p['estagio']} | {p['sem']:.0f} / {p['com']:.0f} | {p['exigido']:.0f} | {br(p['taxa_sem'])} | "
                      f"{br(p['taxa_com'])} (+{p['nivel']} pontos) | ≤ {br(META_SEM_TREINO, 1)} / ≥ {br(META_COM_TREINO, 1)} |")
    linhas += [
        "",
        "Poder = total de stats × multiplicador de estágio (modelo PROPOSTO da planilha). Fontes: `digilab/config.py` "
        "(ConfigCampanha, ConfigCatalogo, ConfigTreino), `experiments/calibrar_progressao.py > ALCANCE` e "
        "`experiments/configs/progressao/validacao_resultado.json` (taxas, sementes de validação).",
        "",
        "Leitura: " + leitura(ps),
        "",
        "## Os 5 alertas mais graves",
        "",
    ]
    for i, (_, t, n, fo, p) in enumerate(graves, start=1):
        linhas += [f"{i}. **{t}.** {n}. Fonte: `{fo}`."]
    linhas += [
        "",
        "## Propostas de ajuste, em ordem de prioridade (só propostas)",
        "",
    ]
    for i, (_, t, n, fo, p) in enumerate(graves, start=1):
        linhas.append(f"{i}. {p}")
    linhas += [
        "",
        "Divergências entre `config.py`, `docs/*.json` e a web e as células SEM FONTE estão na aba `alertas`.",
        "",
    ]
    SAIDA_MD.write_text("\n".join(linhas), encoding="utf-8")


# ------------------------------------------------------------------ montagem e verificacao

def montar(destino: Path, com_png: bool = True) -> tuple[Workbook, list[Aba], dict]:
    f = Fontes()
    wb = Workbook()
    aba_leia_me(wb, f)
    abas = [aba_referencias(wb), aba_jornada(wb, f), aba_estagios(wb, f), aba_portoes(wb, f), aba_economia(wb, f),
            aba_endgame(wb), aba_principios(wb), aba_escala(wb, f), aba_alertas(wb, f)]
    fixo = datetime.fromisoformat(f.data_commit).replace(tzinfo=None)
    wb.properties.created = fixo
    wb.properties.modified = fixo
    wb.properties.creator = "tools/gerar_planilha_progressao.py"
    wb.save(destino)
    curva = curva_png(f, SAIDA_PNG) if com_png else {}
    return wb, abas, {"fontes": f, "curva": curva}


def valores(caminho: Path) -> dict:
    wb = load_workbook(caminho)
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


def verificar() -> None:
    """Gera duas vezes em pastas temporarias: mesmas celulas; toda entrada numerica com fonte e status."""
    with tempfile.TemporaryDirectory() as d:
        p1, p2 = Path(d) / "a.xlsx", Path(d) / "b.xlsx"
        _, abas, _ = montar(p1, com_png=False)
        montar(p2, com_png=False)
        if valores(p1) != valores(p2):
            raise SystemExit("FALHA: duas geracoes com valores diferentes")
        wb = load_workbook(p1)
        n = 0
        for a in abas:
            ws = wb[a.ws.title]
            for ref in a.entradas:
                c = ws[ref]
                fonte, status = ws.cell(c.row, c.column + 1).value, ws.cell(c.row, c.column + 2).value
                if status not in STATUS or (status != "SEM FONTE" and not fonte):
                    raise SystemExit(f"FALHA: {ws.title}!{ref} sem fonte/status")
                n += 1
            # nenhum numero solto fora dos trios de entrada (formulas sao texto que comeca com '=')
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, (int, float)) and not isinstance(c.value, bool) and c.coordinate not in a.entradas \
                            and ws.title not in ("alertas",):
                        raise SystemExit(f"FALHA: {ws.title}!{c.coordinate} = {c.value} e numero sem fonte")
        print(f"ok: duas geracoes identicas; {n} entradas numericas, todas com fonte e status")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--verificar", action="store_true")
    a = ap.parse_args()
    if a.verificar:
        verificar()
        return
    SAIDA_DIR.mkdir(exist_ok=True)
    _, abas, extra = montar(SAIDA_XLSX)
    escrever_md(extra["fontes"], extra["curva"])
    sem = sum(1 for ab in abas for ref in ab.entradas if ab.ws[ref].value == "SEM FONTE")
    print(f"planilha: {SAIDA_XLSX.relative_to(RAIZ)}; curva: {SAIDA_PNG.relative_to(RAIZ)}; resumo: {SAIDA_MD.relative_to(RAIZ)}")
    print(f"entradas: {sum(len(ab.entradas) for ab in abas)}; SEM FONTE: {sem}")


if __name__ == "__main__":
    main()
